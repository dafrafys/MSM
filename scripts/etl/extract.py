# scripts/etl/extract.py
import os
import pandas as pd
import re
from pathlib import Path
from openpyxl import load_workbook
from datetime import date, datetime
from scripts.etl.common import flatten_multiindex, log_info, get_path, normaliser_nom


def extraire_records_de_fichier(path: Path, colonnes=(0, 1, 2)):
    """
    Lit un fichier Excel (BBS) et renvoie une liste de dicts {Personne, Date_BBS, Description}.
    - path : objet pathlib.Path pointant vers un .xlsx
    - colonnes : indices (0=col date, 1=col observateur, 2=col tâche/description)
    """
    try:
        wb = load_workbook(path, data_only=True)
    except PermissionError:
        log_info(f"⚠️ Impossible d’ouvrir {path.name}, fichier ignoré")
        return []

    ws = wb[wb.sheetnames[0]]
    recs = []
    for row in ws.iter_rows(min_row=4, max_col=max(colonnes) + 1):
        cell_date   = row[colonnes[0]].value
        observateur = row[colonnes[1]].value
        tache       = row[colonnes[2]].value

        # Si pas d’observateur ou pas de date, on ignore
        if not observateur or not cell_date:
            continue

        # Si la cellule date est un datetime (vraie date Excel)
        if isinstance(cell_date, datetime):
            d = cell_date.date()
            recs.append({
                "Personne":    str(observateur).strip(),
                "Date_BBS":    d,
                "Description": str(tache).strip()
            })
        else:
            # Si c’est un texte du type "5; 12; 19", on extrait tous les nombres
            for j in re.findall(r"\d{1,2}", str(cell_date)):
                try:
                    d = date(date.today().year, date.today().month, int(j))
                    recs.append({
                        "Personne":    str(observateur).strip(),
                        "Date_BBS":    d,
                        "Description": str(tache).strip()
                    })
                except:
                    continue

    return recs


def extract_bbs(fallback_dir: str = None) -> pd.DataFrame:
    """
    Parcourt le dossier BBS et renvoie un DataFrame brut contenant
    toutes les lignes trouvées sous forme de colonnes [Personne, Date_BBS, Description].
    - fallback_dir : chemin local (data/) où copier un ou plusieurs fichiers si
                     le lecteur réseau n’est pas monté en dev.
    """
    log_info("Extract (BBS) : début de l'extraction des fichiers Excel")

    # 1) Dossiers à scanner (en production, chemin réseau)
    dossiers = [Path(get_path("BBS_PATH"))]

    # 2) Lister tous les .xlsx valides (ignorer les temp ~*)
    fichiers = []
    for d in dossiers:
        for f in d.glob("*.xlsx"):
            if not f.name.startswith("~$"):
                fichiers.append(f)

    # 3) Si aucun fichier trouvé et qu’on a un fallback_dir, on tente le dossier local dev
    if not fichiers and fallback_dir:
        log_info(f"Extract (BBS) : aucun fichier sur réseau, tentative dans {fallback_dir}")
        local_path = Path(fallback_dir)
        if local_path.exists():
            for f in local_path.glob("*.xlsx"):
                if not f.name.startswith("~$"):
                    fichiers.append(f)

    if not fichiers:
        raise FileNotFoundError("Extract (BBS) : aucun fichier .xlsx trouvé ni en production ni en fallback")

    # 4) Parcourir chaque fichier et accumuler les enregistrements
    all_records = []
    for f in fichiers:
        log_info(f"Extract (BBS) : lecture de {f.name}")
        recs = extraire_records_de_fichier(f)
        all_records.extend(recs)

    df = pd.DataFrame(all_records)
    if df.empty:
        log_info("Extract (BBS) : aucun enregistrement trouvé dans les fichiers")
    else:
        log_info(f"Extract (BBS) : {len(df)} enregistrement(s) brut(s) collecté(s)")

    return df

def extract_accidents(filepath: str = None) -> pd.DataFrame:
    """
    Lit le fichier Excel 'ACCI 2025 integration DU.xlsm'.
    - Si filepath n'est pas fourni, on prend le chemin d'origine sur le lecteur réseau.
    - En local, on peut copier l'Excel dans data/ACCI 2025 integration DU.xlsm pour un fallback.
    """
    log_info("Extract (Accidents) : début de la lecture du fichier Excel")

    # Chemin par défaut vers l'Excel sur le lecteur réseau
    #default_path = "/mnt/pdrivedrivedrive/1 - EHS/04-SECURITE/01-ACCIDENTS/2025/ACCI 2025 integration DU.xlsm"
    if filepath is None:
        filepath = get_path("ACCIDENT_XLSM")

    # Si le fichier n'existe pas au chemin réseau, on essaie un fallback dans data/
    if not os.path.exists(filepath):
        log_info(f"Extract (Accidents) : fichier non trouvé à {filepath}, tentative dans data/")
        base = os.path.abspath(os.path.join(__file__, "../../data"))
        fallback = os.path.join(base, "ACCI 2025 integration DU.xlsm")
        if os.path.exists(fallback):
            filepath = fallback
            log_info(f"Extract (Accidents) : lecture depuis fallback {fallback}")
        else:
            raise FileNotFoundError(f"Fichier introuvable ni à {filepath} ni en fallback ({fallback})")

    # Lecture du contenu du fichier Excel
    df = pd.read_excel(
        filepath,
        sheet_name="FYSOL ACCIDENT",
        skiprows=4,
        header=0,
        usecols=[
            "N° Accident",
            "Type AT",
            "Atelier",
            "Position/\n N° machine",
            "Date\nAnnée",
            "Date\nMois",
            "Date\nJour",
            "Heure",
            "Unité de travail DU\nà laquelle l'employé est rattachée",
            "Circonstances",
            "DU Classes de risques",
            "Etat Analyse"
        ],
        engine="openpyxl"
    )

    # Renommer les colonnes pour retirer les sauts de ligne
    df.columns = df.columns.str.replace("\n", " ")

    log_info(f"Extract (Accidents) : {len(df)} lignes lues depuis {filepath}")
    return df


def extract_impact_fibrage(filepath: str = None) -> pd.DataFrame:
    """
    Lit le fichier Excel '2025_Suivi des pertes fibrage.xlsx' à son emplacement d'origine
    si filepath n'est pas fourni.
    - En production, vous pouvez conserver le chemin absolu vers SharePoint/lecteur réseau.
    - En dev, vous pouvez copier le fichier dans data/ et le laisser vide pour fallback.
    """
    log_info("Extract (Impact Fibrage) : début de la lecture du fichier Excel")

    # 1) Valeur par défaut = chemin d'origine sur le lecteur réseau SharePoint
    #    Si vous déployez en prod, laissez ce chemin en dur ou dans une variable d'env.
    # default_path = "/mnt/pdrivedrivedrive/3 - Production/1 - Quotidien/2025_Suivi des pertes fibrage.xlsx"
    if filepath is None:
        filepath = get_path("FIBRAGE_XLSX")
    # 2) Si le fichier n'existe pas à ce chemin (ex. en local dev),
    #    on tente de lire depuis data/ (fallback)
    if not os.path.exists(filepath):
        log_info(f"Extract (Impact Fibrage) : fichier non trouvé à {filepath}, tentative dans data/")
        base = os.path.abspath(os.path.join(__file__, "../../data"))
        fallback = os.path.join(base, "2025_Suivi des pertes fibrage.xlsx")
        if os.path.exists(fallback):
            filepath = fallback
            log_info(f"Extract (Impact Fibrage) : lecture depuis fallback {fallback}")
        else:
            raise FileNotFoundError(f"Fichier introuvable ni à {filepath} ni en fallback ({fallback})")

    # 3) Lire le contenu du fichier, avec header sur deux lignes
    df = pd.read_excel(
        filepath,
        sheet_name="Suivi des pertes",
        skiprows=4,
        header=[0, 1],
        engine="openpyxl"
    )

    log_info(f"Extract (Impact Fibrage) : {len(df)} lignes lues depuis {filepath}")
    return df

def extract_impact_finissage(excel_arg: str = None) -> pd.DataFrame:
    """
    Lit le fichier Excel d'Impact Finissage:
    - excel_arg peut être un chemin de dossier ou de fichier.
    - Si dossier, cherche un .xlsx valide dans ce dossier.
    - Sinon, utilise le chemin par défaut en prod.
    Retourne un DataFrame brut avec colonnes multi-index conservées.
    """
    log_info("Extract (Impact Finissage) : détermination du chemin de l'Excel")
    # Déterminer le chemin
    if excel_arg:
        if os.path.isdir(excel_arg):
            # chercher dans le dossier
            for f in Path(excel_arg).glob("*.xlsx"):
                if not f.name.startswith("~$"):
                    excel_path = str(f)
                    log_info(f"Extract : trouvé {excel_path} dans dossier fourni")
                    break
            else:
                raise FileNotFoundError(f"Aucun .xlsx valide dans {excel_arg}")
        elif os.path.isfile(excel_arg):
            excel_path = excel_arg
            log_info(f"Extract : chemin spécifié {excel_path}")
        else:
            raise FileNotFoundError(f"Chemin invalide: {excel_arg}")
    else:
        excel_path = os.path.join(get_path("FINISSAGE_XLSX"))
        log_info(f"Extract : chemin par défaut {excel_path}")

    # Lecture du fichier avec header multi-index
    df = pd.read_excel(
        excel_path,
        sheet_name="Suivi des pertes",
        skiprows=1,
        header=[0,1],
        engine="openpyxl"
    )
    # Affiche les colonnes pour vérifier leur nom
    print("Colonnes dans le fichier Excel :")
    print(df.columns.tolist())
    log_info(f"Extract : {df.shape[0]} lignes x {df.shape[1]} colonnes lues")
    return df
