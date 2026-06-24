import sys
import time
import os
import pandas as pd
from openpyxl import load_workbook
import logging

"""
update_fibrage.py
-----------------
Ce script prend un impact_id, extrait le plan d'action associé depuis la base PostgreSQL,
puis met à jour le fichier Excel de suivi à la bonne ligne.
Il s’appuie sur les utilitaires du module 'common.py' pour :
- se connecter à la base (via la variable d’environnement SQLALCHEMY_DATABASE_URI)
- logguer les infos principales

Test ? À utiliser ainsi :
    python update_fibrage.py <impact_id>
"""

# *** IMPORTS DES UTILITAIRES COMMUNS ***
sys.path.append(os.path.dirname(__file__))  # Pour import depuis le même dossier 'etl'
import common

# 1) --- LOGGING ---
logging.basicConfig(
    level=logging.DEBUG,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S"
)

# 2) --- RÉCUPÉRATION DE L’ID DE L’IMPACT ---
if len(sys.argv) < 2:
    logging.error("Usage: python update_fibrage.py <impact_id>")
    sys.exit(1)
impact_id = sys.argv[1]

# 3) --- CONNEXION À LA BASE ---
try:
    engine = common.get_sqlalchemy_engine()
except Exception as e:
    logging.error("Erreur connexion base : %s", e)
    sys.exit(1)

# 4) --- EXTRACTION DU PLAN D’ACTION ---
sql = """
    SELECT p.*, i."Excel_row"
    FROM plan_daction p
    JOIN "Impacts" i ON p."Impact_Id" = i."Impact_Id"
    WHERE i."Impact_Id" = :impact_id
"""
df = pd.read_sql(sql, engine, params={"impact_id": impact_id})

if df.empty:
    common.log_info(f"Aucune donnée à exporter pour l'impact_id {impact_id}")
    print("Pas de données à exporter")
    sys.exit(0)

# 5) --- CHEMINS ET PARAMS ---
EXCEL_PATH = common.get_path("FIBRAGE_XLSX")
EXCEL_SHEET = "Suivi des pertes"
MAX_ATTEMPTS = 2
DELAY_SEC    = 10 * 60

col_mapping = {
    "Cause": "AI",
    "Action": "AJ",
    "Commentaire_At": "AK",
    "pilot_id": "AL",
    "Echeance_at": "AM",
    "Statut": "AN"
}

def export_plan_action_to_excel(df, excel_path, sheet_name, col_mapping):
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            logging.debug("Tentative %d – Chargement du fichier Excel", attempt)
            wb = load_workbook(excel_path)
            ws = wb[sheet_name]
            existing_dvs = list(ws.data_validations.dataValidation)
            logging.info("Validations initiales (avant écriture) : %d objets", len(existing_dvs))
            for _, row in df.iterrows():
                row_idx = int(row["Excel_row"])
                for col_name, excel_col in col_mapping.items():
                    val = row.get(col_name)
                    if pd.notna(val):
                        ws[f"{excel_col}{row_idx}"] = str(val)
            ws.data_validations.dataValidation = []
            for dv in existing_dvs:
                ws.add_data_validation(dv)
            logging.info("Validations réinjectées : %d objets", len(existing_dvs))
            wb.save(excel_path)
            logging.debug("Fichier Excel mis à jour et sauvegardé.")
            print("OK Plan d’action exporté vers Excel.")
            return True
        except PermissionError:
            logging.warning("Fichier Excel ouvert par un autre utilisateur. Tentative %d/%d", attempt, MAX_ATTEMPTS)
            if attempt < MAX_ATTEMPTS:
                logging.info("Nouvelle tentative dans %d secondes...", DELAY_SEC)
                time.sleep(DELAY_SEC)
            else:
                logging.error("Impossible d'écrire après %d tentatives.", MAX_ATTEMPTS)
                print("Échec définitif après 2 tentatives. Réessayer plus tard.")
                return False
        except Exception as e:
            logging.error("Erreur inattendue lors de l'export Excel : %s", e)
            print(f"Erreur inattendue : {e}")
            return False

if __name__ == "__main__":
    export_plan_action_to_excel(df, EXCEL_PATH, EXCEL_SHEET, col_mapping)
